# -*- coding: utf-8 -*-
"""Genera el diccionario de datos y la matriz de correspondencia del modelo final (Sprint 5).

Fuentes (todas versionadas en el repositorio, sin datos de clientes):
  etl/adn_etl/modelo.py                         tablas y columnas del modelo (fuente de verdad del DDL)
  etl/adn_etl/descripciones.py                  descripciones de las columnas que modelo.py deja vacías
  etl/adn_etl/rag.py                            disponibilidad para ask.data.cr / RAG (decisión D6)
  etl/adn_etl/datos/matriz_correspondencia.csv  matriz aprobada que siguen los pasos de carga
  etl/adn_etl/datos/matriz_perfilado_sprint1.csv perfilado del Sprint 1 (columnas informativas de la matriz)
  sql/01_control/01_control_etl.sql             tablas de control agregadas en los Sprints 3 y 4 (CONTROL_EXTRA)

Uso:  python herramientas/gen_diccionario.py [salida.xlsx]
Después se recalculan las fórmulas al abrir el archivo en Excel (o con LibreOffice).
"""
import csv
import os
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

RAIZ = os.path.join(os.path.dirname(__file__), "..")
sys.path.insert(0, os.path.join(RAIZ, "etl"))
sys.path.insert(0, os.path.dirname(__file__))
from adn_etl.modelo import T, LINAJE  # noqa: E402
from adn_etl.descripciones import describir  # noqa: E402
from adn_etl import rag  # noqa: E402
from gen_ddl import REGLAS  # noqa: E402

VERSION = "2.0"
FECHA = "26/10/2026"
DATOS = os.path.join(RAIZ, "etl", "adn_etl", "datos")

AZUL = "1F3864"
F_TIT = Font(name="Arial", size=14, bold=True, color=AZUL)
F_SUB = Font(name="Arial", size=10)
F_SEC = Font(name="Arial", size=11, bold=True)
F_HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TXT = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
FILL = PatternFill("solid", fgColor=AZUL)
ARRIBA = Alignment(vertical="top")
AJUSTE = Alignment(vertical="top", wrap_text=True)

# Tablas de control creadas por sql/01_control/01_control_etl.sql (no están en modelo.py)
_N, _V = (lambda n: f"nvarchar({n})"), (lambda n: f"varchar({n})")
DT = "datetimeoffset(0)"
CONTROL_EXTRA = [
    ("Bitacora", "Bitácora por registro de cada carga confirmada; permite el rollback por lote", "Una fila por registro insertado, actualizado o retirado", [
        ("BitacoraId", "bigint", False, "Llave sustituta", "", True, None),
        ("LoteId", "int", False, "Lote que hizo el cambio", "IN-05", False, "ctl.Lote"),
        ("TablaModelo", _N(128), False, "Tabla del modelo afectada", "", False, None),
        ("RegistroId", "bigint", False, "Llave sustituta del registro afectado", "", False, None),
        ("Accion", "char(1)", False, "I = insertado, U = actualizado, R = retirado porque el origen ya no lo produce", "", False, None),
        ("FuenteLlave", _N(200), True, "Llave del registro en el origen", "IN-05", False, None),
        ("ValoresAnteriores", "nvarchar(max)", True, "JSON con los valores antes del cambio (U y R); el rollback los restaura", "", False, None),
        ("Fecha", DT, False, "Fecha y hora del cambio (hora de Costa Rica)", "ES-01", False, None)]),
    ("Entorno", "Marca de entorno: el ETL se niega a ejecutar si la base no tiene la marca SANDBOX", "Una fila", [
        ("Nombre", _V(20), False, "Nombre del entorno (SANDBOX)", "", True, None),
        ("FechaRegistro", DT, False, "Fecha en que se creó la marca", "", False, None),
        ("Nota", _N(400), True, "Advertencia: nunca crear esta marca en producción", "", False, None)]),
    ("Reconciliacion", "Resultado de ./etl.ps1 reconciliar por tabla de origen (Sprint 4)", "Una fila por tabla de origen y ejecución", [
        ("ReconciliacionId", "bigint", False, "Llave sustituta", "", True, None),
        ("Ejecucion", DT, False, "Fecha y hora de la ejecución", "", False, None),
        ("Paso", _N(150), False, "Paso de carga que lee la tabla de origen", "", False, None),
        ("TablaModelo", _N(128), False, "Tabla destino", "", False, None),
        ("FuenteBase", _N(128), False, "Base de origen", "", False, None),
        ("FuenteTabla", _N(128), False, "Tabla de origen", "", False, None),
        ("FilasOrigen", "bigint", True, "Filas de la tabla de origen", "", False, None),
        ("FueraDeAlcance", "bigint", False, "Filas excluidas por un filtro de alcance documentado", "", False, None),
        ("FilasUniverso", "bigint", False, "Filas que la carga debe considerar", "", False, None),
        ("Repetidas", "bigint", False, "Filas con llave repetida", "DD-04", False, None),
        ("LlaveVacia", "bigint", False, "Filas sin llave", "", False, None),
        ("LlaveVaciaJustificada", "bigint", False, "Filas sin llave con motivo en ctl.Descarte", "", False, None),
        ("Llaves", "bigint", False, "Llaves distintas", "", False, None),
        ("Cargadas", "bigint", False, "Llaves presentes en el modelo", "", False, None),
        ("Excluidas", "bigint", False, "Llaves en ctl.Descarte de un lote confirmado", "DD-06", False, None),
        ("SinConciliar", "bigint", False, "Llaves sin destino ni motivo (debe ser 0)", "", False, None),
        ("DestinoSinOrigen", "bigint", False, "Registros del modelo que el origen ya no produce (debe ser 0)", "", False, None),
        ("FilasConciliadas", "bigint", False, "Filas con destino o motivo", "", False, None),
        ("PctConciliado", "decimal(9,4)", False, "Porcentaje conciliado", "", False, None),
        ("Estado", _V(10), False, "OK o REVISAR", "", False, None),
        ("Segundos", "decimal(9,1)", True, "Duración de la consulta", "", False, None)]),
    ("ReconciliacionPendiente", "Llaves técnicas sin conciliar o en el modelo sin origen (Sprint 4)", "Una fila por llave pendiente", [
        ("PendienteId", "bigint", False, "Llave sustituta", "", True, None),
        ("Ejecucion", DT, False, "Ejecución de la reconciliación", "", False, None),
        ("Paso", _N(150), False, "Paso de carga", "", False, None),
        ("FuenteBase", _N(128), False, "Base de origen", "", False, None),
        ("FuenteTabla", _N(128), False, "Tabla de origen", "", False, None),
        ("FuenteLlave", _N(200), False, "Llave pendiente", "", False, None),
        ("Tipo", _V(20), False, "SinConciliar o DestinoSinOrigen", "", False, None)]),
]
LOTE_EXTRA = [  # columnas que 01_control_etl.sql agrega a ctl.Lote
    ("Paso", _N(100), "Paso de carga ejecutado"), ("Modo", _V(10), "dry-run (ensayo con ROLLBACK) o confirm"),
    ("FilasActualizadas", "bigint", "Filas actualizadas"), ("FilasSinCambio", "bigint", "Filas que no cambiaron (mismo HashContenido)"),
    ("Anomalias", "bigint", "Anomalías registradas en ctl.Anomalia"), ("Parametros", _N(400), "Parámetros de la ejecución (p. ej. período)"),
    ("ArchivoRollback", _N(400), "Archivo con el script de rollback del lote"), ("Mensaje", _N(2000), "Mensaje de error o resumen"),
    ("Equipo", _N(128), "Equipo donde se ejecutó"), ("FilasExcluidas", "bigint", "Filas leídas menos válidas (conciliación)"),
    ("LlavesSinMotivo", "bigint", "Llaves excluidas sin registro en ctl.Descarte (debe ser 0)"),
]

DISPOSICIONES = ["Mapeada", "Atributo extendido", "Uso en validación", "Filtro de carga", "No se migra",
                 "Excluida (IN-04)", "Tabla descartada"]
BASES = ["ADNApp", "ADNPayV2", "Navsat", "ZBooks_Mirror", "ZCRM2Trengo", "Zoho_Mirror", "History_Tables", "app.data.cr (MongoDB)"]


def tipo_txt(c):
    return "char(7) calculada" if c["type"] == "computed" else c["type"]


def leer_csv(nombre):
    with open(os.path.join(DATOS, nombre), encoding="utf-8") as f:
        return list(csv.DictReader(f))


def hoja(wb, titulo, cabecera, filas, anchos, ajustar=(), nombre_tabla=None, formatos=None):
    ws = wb.create_sheet(titulo)
    ws.append(cabecera)
    for i, _ in enumerate(cabecera, 1):
        c = ws.cell(1, i)
        c.font, c.fill, c.alignment = F_HDR, FILL, Alignment(vertical="center", wrap_text=True)
    for f in filas:
        ws.append(list(f))
    for col, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.font = F_TXT
            c.alignment = AJUSTE if c.column in ajustar else ARRIBA
            if formatos and c.column in formatos:
                c.number_format = formatos[c.column]
    ws.freeze_panes = "A2"
    ref = f"A1:{get_column_letter(len(cabecera))}{max(ws.max_row, 2)}"
    ws.add_table(Table(displayName=nombre_tabla or titulo, ref=ref))
    return ws


def destino_tipos(tabla_dest, col_dest, cols_modelo, v1):
    """Tipo de las columnas destino de una fila de la matriz ('a; b' → 'tipo a; tipo b')."""
    if not tabla_dest or not col_dest:
        return ""
    tipos = []
    for parte in col_dest.split(";"):
        m = re.match(r"\s*([A-Za-z0-9_]+(?:\.[A-Za-z0-9_]+)*)", parte)
        if not m:
            continue
        nombre = m.group(1)
        if nombre.startswith("AtributosExtendidos"):
            tipos.append("nvarchar(max) JSON")
            continue
        if nombre.count(".") == 2:            # core.Cliente.Telefono
            t, c = nombre.rsplit(".", 1)
        else:
            t, c = tabla_dest, nombre
        tp = cols_modelo.get((t, c))
        if tp:
            tipos.append(tp)
    return "; ".join(tipos) if tipos else (v1 or "")


def construir(salida):
    # ---------------- modelo: tablas y columnas ----------------
    matriz = leer_csv("matriz_correspondencia.csv")
    perf = {(r["Base"], r["TablaOrigen"], r["ColumnaOrigen"]): r for r in leer_csv("matriz_perfilado_sprint1.csv")}

    origen_de = defaultdict(list)  # (tabla, columna) -> ['Base.Tabla.Columna']
    for r in matriz:
        if r["Disposicion"] != "Mapeada" or not r["TablaDestino"]:
            continue
        for parte in r["ColumnaDestino"].split(";"):
            m = re.match(r"\s*([A-Za-z0-9_]+(?:\.[A-Za-z0-9_]+)*)", parte)
            if not m:
                continue
            nombre = m.group(1)
            t, c = (nombre.rsplit(".", 1) if nombre.count(".") == 2 else (r["TablaDestino"], nombre))
            ref = f"{r['Base']}.{r['TablaOrigen']}.{r['ColumnaOrigen']}"
            if ref not in origen_de[(t, c)]:
                origen_de[(t, c)].append(ref)

    tablas, dic, cols_modelo = [], [], {}
    for t in T:
        nombre = f"{t['schema']}.{t['name']}"
        cols = list(t["cols"]) + (list(LINAJE) if t["linaje"] else [])
        if nombre == "ctl.Lote":
            cols += [dict(name=n, type=tp, null=True, desc=d, regla="") for n, tp, d in LOTE_EXTRA]
        for orden, c in enumerate(cols, 1):
            cols_modelo[(nombre, c["name"])] = tipo_txt(c)
            es_lin = any(c["name"] == x["name"] for x in LINAJE) and t["linaje"]
            if c.get("identity"):
                org = "Generada (IDENTITY)"
            elif es_lin:
                org = "Generada en la carga (trazabilidad)"
            elif c["name"] == "AtributosExtendidos":
                org = "Columnas de origen marcadas como Atributo extendido en la matriz"
            else:
                org = "; ".join(origen_de.get((nombre, c["name"]), []))
            cat, motivo = rag.columna(t["schema"], t["name"], c["name"])
            dic.append([t["schema"], t["name"], t["area"], orden, c["name"], tipo_txt(c), "Sí" if c["null"] else "No",
                        "Sí" if c.get("pk") or t.get("pkcol") == c["name"] else None,
                        c.get("fk") or ("ctl.Lote" if es_lin and c["name"] == "LoteId" else None),
                        describir(t["name"], c["name"], c["desc"]), c["regla"] or None, "Sí" if es_lin else None,
                        org or None, cat, motivo or None])
        nivel, motivo = rag.tabla(t["schema"], t["name"], [c["name"] for c in cols])
        tablas.append([t["schema"], t["name"], t["area"], t["desc"], t["grano"], t["origen"] or None, len(cols),
                       "Sí" if t["linaje"] else "Control", nivel, motivo or None])
    for nombre, desc, grano, cols in CONTROL_EXTRA:
        for orden, (cn, tp, nul, d, regla, es_pk, fk) in enumerate(cols, 1):
            cat, motivo = rag.columna("ctl", nombre, cn)
            dic.append(["ctl", nombre, "Control", orden, cn, tp, "Sí" if nul else "No", "Sí" if es_pk else None, fk, d,
                        regla or None, None, None, cat, motivo or None])
        nivel, motivo = rag.tabla("ctl", nombre, [c[0] for c in cols])
        tablas.append(["ctl", nombre, "Control", desc, grano, "sql/01_control/01_control_etl.sql", len(cols), "Control",
                       nivel, motivo or None])

    # ---------------- matriz ----------------
    mat = []
    for i, r in enumerate(matriz, 1):
        p = perf.get((r["Base"], r["TablaOrigen"], r["ColumnaOrigen"]), {})
        pct = p.get("PctVacios", "")
        mat.append([i, r["Sistema"], r["Base"], r["TablaOrigen"], r["ColumnaOrigen"], r["TipoOrigen"],
                    r["LlavePrimaria"] or None, p.get("Perfilado") or None, float(pct) if pct not in ("", None) else None,
                    r["Disposicion"], r["TablaDestino"] or None, r["ColumnaDestino"] or None,
                    destino_tipos(r["TablaDestino"], r["ColumnaDestino"], cols_modelo, p.get("TipoDestinoV1")) or None,
                    r["Regla"] or None, r["Observacion"] or None])

    wb = Workbook()
    wb.remove(wb.active)
    res = wb.create_sheet("Resumen")

    hoja(wb, "Tablas", ["Esquema", "Tabla", "Área", "Descripción", "Grano", "Origen principal", "Columnas",
                        "Trazabilidad IN-05", "ask.data.cr / RAG", "Motivo (RAG)"], tablas,
         [9, 26, 18, 60, 36, 55, 10, 14, 26, 60], ajustar=(4, 5, 6, 10), nombre_tabla="TablasModelo")
    hoja(wb, "Diccionario", ["Esquema", "Tabla", "Área", "Orden", "Columna", "Tipo", "Admite nulos", "PK",
                             "FK (referencia)", "Descripción", "Regla", "Trazabilidad", "Origen", "ask.data.cr / RAG",
                             "Motivo (RAG)"], dic,
         [8, 24, 16, 6, 28, 18, 8, 5, 20, 55, 9, 11, 60, 22, 45], ajustar=(10, 13, 15), nombre_tabla="Diccionario")
    hoja(wb, "Matriz", ["N", "Sistema", "Base", "Tabla origen", "Columna origen", "Tipo origen", "PK",
                        "Perfilado Sprint 1", "% vacíos", "Disposición", "Tabla destino", "Columna destino",
                        "Tipo destino", "Regla", "Observación"], mat,
         [6, 11, 18, 26, 28, 14, 5, 13, 8, 17, 24, 40, 18, 14, 55], ajustar=(12, 15), nombre_tabla="Matriz",
         formatos={9: "0.0%"})
    hoja(wb, "Trazabilidad", ["Columna o tabla", "Tipo", "Admite nulos", "Descripción", "Regla"],
         [[c["name"], tipo_txt(c), "Sí" if c["null"] else "No", c["desc"], c["regla"]] for c in LINAJE] + [
             ["ctl.LinajeCampo", "tabla", None, "En Cliente, Servicio y Producto (consolidados de varios sistemas) guarda qué sistema, tabla, llave y fecha aportó cada campo ganador según D9; desde v1.3 por lote", "IN-05, D9"],
             ["core.ClienteOrigen", "tabla", None, "Tabla de correspondencia: cada ID de cada sistema, el cliente/servicio al que se vinculó y el método (circuito, cédula, sin correspondencia)", "ID-04, ID-05"],
             ["ctl.Lote", "tabla", None, "Cada ejecución de carga: paso, modo, inicio, fin, filas leídas/cargadas/actualizadas/descartadas, anomalías y archivo de rollback", "IN-05"],
             ["ctl.Bitacora", "tabla", None, "Cada registro insertado, actualizado o retirado por un lote confirmado, con sus valores anteriores; base del rollback", "IN-05"],
             ["ctl.Descarte", "tabla", None, "Registros no cargados con la regla y la llave conservada (duplicados, eliminados, sin referencia obligatoria)", "DD-06"],
             ["ctl.Anomalia", "tabla", None, "Valores marcados: fechas ficticias o futuras, cédulas inválidas, huérfanos, textos recortados", "ES-02, ES-03, IN-01"]],
         [24, 18, 10, 100, 14], ajustar=(4,), nombre_tabla="Trazabilidad")
    hoja(wb, "Reglas", ["Regla", "Tipo", "Descripción", "Criticidad", "Opción aprobada", "ask.data.cr / RAG"],
         [[r[0], r[1], r[2], r[3], r[4], "Disponible"] for r in REGLAS], [9, 16, 100, 11, 10, 18], ajustar=(3,),
         nombre_tabla="Reglas")
    hoja(wb, "Cambios", ["Versión", "Fecha", "Sprint", "Cambio", "Script o commit", "Aprobó"], CAMBIOS,
         [9, 11, 9, 100, 40, 18], ajustar=(4, 5), nombre_tabla="Cambios")
    hoja(wb, "Pendientes", ["Id", "Tema", "Detalle (Sprint 2)", "Responsable", "Resolución", "Estado"], PENDIENTES,
         [7, 34, 70, 22, 80, 12], ajustar=(3, 5), nombre_tabla="Pendientes")

    # ---------------- Resumen (fórmulas sobre las demás hojas) ----------------
    res["A1"] = "Matriz de correspondencia y diccionario de datos · Modelo unificado ADN_Unificado"
    res["A1"].font = F_TIT
    res["A2"] = (f"Sprint 5 (cierre) · American Data Networks S.A. · Luis Felipe Díaz Guadamuz · Versión {VERSION} "
                 f"({FECHA}) · modelo final, generado con herramientas/gen_diccionario.py")
    res["A2"].font = F_SUB
    res["A4"] = "Cobertura por base de origen (fórmulas sobre la hoja Matriz)"
    res["A4"].font = F_SEC
    cab = ["Base de origen", "Columnas"] + DISPOSICIONES + ["Sin disposición", "% con disposición"]
    for i, v in enumerate(cab, 1):
        c = res.cell(5, i, v)
        c.font, c.fill, c.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), FILL, Alignment(wrap_text=True, vertical="center")
    for k, base in enumerate(BASES):
        f = 6 + k
        res.cell(f, 1, base)
        res.cell(f, 2, f"=COUNTIFS(Matriz!$C:$C,$A{f})")
        for j in range(len(DISPOSICIONES)):
            col = get_column_letter(3 + j)
            res.cell(f, 3 + j, f"=COUNTIFS(Matriz!$C:$C,$A{f},Matriz!$J:$J,{col}$5)")
        res.cell(f, 10, f"=B{f}-SUM(C{f}:I{f})")
        res.cell(f, 11, f"=IF(B{f}=0,0,SUM(C{f}:I{f})/B{f})")
    ft = 6 + len(BASES)
    res.cell(ft, 1, "Total")
    for j in range(2, 11):
        col = get_column_letter(j)
        res.cell(ft, j, f"=SUM({col}6:{col}{ft - 1})")
    res.cell(ft, 11, f"=IF(B{ft}=0,0,SUM(C{ft}:I{ft})/B{ft})")
    for f in range(6, ft + 1):
        for j in range(1, 12):
            c = res.cell(f, j)
            c.font = Font(name="Arial", size=9, bold=(f == ft))
            c.number_format = "0.0%" if j == 11 else "#,##0"

    f0 = ft + 2
    res.cell(f0, 1, "Modelo destino").font = F_SEC
    filas_modelo = [
        ("Tablas del modelo (core, ref, ctl)", "=COUNTA(Tablas!B:B)-1"),
        ("  core", '=COUNTIFS(Tablas!A:A,"core")'), ("  ref", '=COUNTIFS(Tablas!A:A,"ref")'),
        ("  ctl", '=COUNTIFS(Tablas!A:A,"ctl")'),
        ("Columnas del modelo (incluye trazabilidad)", "=COUNTA(Diccionario!E:E)-1"),
        ("Columnas de trazabilidad por tabla", len(LINAJE)),
        ("Tablas con trazabilidad completa", '=COUNTIFS(Tablas!H:H,"Sí")'),
        ("Llaves foráneas declaradas", '=COUNTIFS(Diccionario!I:I,"?*")-1'),
        ("Columnas de origen evaluadas (SQL Server + MongoDB)", "=COUNTA(Matriz!E:E)-1"),
        ("Reglas y decisiones del catálogo", "=COUNTA(Reglas!A:A)-1"),
    ]
    for k, (txt, val) in enumerate(filas_modelo, 1):
        res.cell(f0 + k, 1, txt).font = Font(name="Arial", size=9)
        c = res.cell(f0 + k, 2, val)
        c.font, c.number_format = Font(name="Arial", size=9), "#,##0"
    res.cell(f0 + 6, 3, "Columnas FuenteSistema … HashContenido (hoja Trazabilidad)").font = Font(name="Arial", size=8, italic=True)

    f1 = f0 + len(filas_modelo) + 2
    res.cell(f1, 1, "Disponibilidad para ask.data.cr y sistemas RAG (decisión D6)").font = F_SEC
    for i, v in enumerate(["Nivel", "Tablas", "", "Categoría de columna", "Columnas", "% del modelo"], 1):
        if v:
            c = res.cell(f1 + 1, i, v)
            c.font, c.fill = Font(name="Arial", size=9, bold=True, color="FFFFFF"), FILL
    niveles = ["Disponible", "Disponible con restricciones", "No disponible"]
    cats = [rag.SI, rag.ANON, rag.PERSONAL, rag.TECNICO]
    for k in range(max(len(niveles), len(cats))):
        f = f1 + 2 + k
        if k < len(niveles):
            res.cell(f, 1, niveles[k])
            res.cell(f, 2, f"=COUNTIFS(Tablas!I:I,A{f})")
        if k < len(cats):
            res.cell(f, 4, cats[k])
            res.cell(f, 5, f"=COUNTIFS(Diccionario!N:N,D{f})")
            res.cell(f, 6, f"=IF(SUM($E${f1 + 2}:$E${f1 + 1 + len(cats)})=0,0,E{f}/SUM($E${f1 + 2}:$E${f1 + 1 + len(cats)}))")
        for j in (1, 2, 4, 5, 6):
            res.cell(f, j).font = Font(name="Arial", size=9)
            res.cell(f, j).number_format = "0.0%" if j == 6 else "#,##0"
    fn = f1 + 3 + len(cats)
    res.cell(fn, 1, ("Disponible = se puede exponer tal cual (vistas del esquema rag, sql/30_rag/vistas_rag.sql). "
                     "Requiere anonimización = texto libre; entra al corpus solo después de anonimizarlo. "
                     "No = dato personal o técnico. Detalle por columna en la hoja Diccionario.")).font = Font(name="Arial", size=8, italic=True)

    f2 = fn + 2
    res.cell(f2, 1, "Cómo leer la columna Disposición").font = F_SEC
    leyenda = [
        ("Mapeada", "La columna tiene una columna propia en el modelo (Columna destino)."),
        ("Atributo extendido", "Se conserva en AtributosExtendidos (JSON) de la tabla destino, con su nombre de origen; no se pierde."),
        ("Uso en validación", "Se usa para validar o derivar otro campo y no se guarda por separado (ES-08)."),
        ("Filtro de carga", "Indica registros eliminados en el origen: esos registros van a ctl.Descarte (DD-06)."),
        ("No se migra", "Vacía, duplicada con formato o metadato de la réplica (IN-03, IN-05)."),
        ("Excluida (IN-04)", "Credencial, dato de tarjeta, foto de cédula o firma: no se copia al modelo."),
        ("Tabla descartada", "Tabla interna del framework, vacía o de prueba; motivo en Observación."),
    ]
    for k, (a, b) in enumerate(leyenda, 1):
        res.cell(f2 + k, 1, a).font = F_B
        res.cell(f2 + k, 2, b).font = F_TXT
    res.column_dimensions["A"].width = 46
    for j in range(2, 12):
        res.column_dimensions[get_column_letter(j)].width = 13
    res.column_dimensions["D"].width = 22
    wb.save(salida)
    return dict(tablas=len(tablas), columnas=len(dic), matriz=len(mat),
                fks=sum(1 for d in dic if d[8]), sin_tipo=sum(1 for m in mat if m[9] == "Mapeada" and not m[12]))


CAMBIOS = [
    ["1.0", "07/10/2026", "2", "Diseño inicial: 74 tablas, matriz de 2.449 columnas de origen con disposición.", "Matriz_Correspondencia_Diccionario_Sprint2.xlsx", "—"],
    ["1.1", "12/10/2026", "2", "Se quita la restricción UNIQUE duplicada de core.Cliente.Cedula (la unicidad la da el índice filtrado UX_Cliente_Cedula, D8). "
     "core.EventoConexion: customer_name trae el ID del contacto de Books (5.350/5.350); se reemplazan NombreCliente y ServicioId por ContactoBooksId y ClienteId. "
     "Telefon_a se mapea como indicador de telefonía (TieneVoIP), no como teléfono.", "sql/00_modelo/02 y 03", "José Hernández"],
    ["1.2", "12/10/2026", "2", "Decisiones D10 (moneda de llamadas según destino), D11 (core.MinutosContratadosHistorico, tabla nueva) y D12 (tickets y conexiones como histórico).", "sql/00_modelo/04", "José Hernández"],
    ["1.3", "12/10/2026", "3", "core.Servicio.Circuito deja de ser único (140 circuitos en más de una cuenta); ctl.LinajeCampo guarda el linaje por lote para que el rollback deje vigente el anterior.", "sql/00_modelo/05 y 06", "José Hernández"],
    ["1.3", "13/10/2026", "3", "Objetos de control del ETL: ctl.Lote ampliado (paso, modo, conteos, archivo de rollback), ctl.Bitacora, ctl.Entorno, funciones de normalización y staging de MongoDB (esquema stg, fuera del modelo).", "sql/01_control/01 y 02", "—"],
    ["1.4", "19/10/2026", "3", "Regla ID-01 v1.1: el circuito conserva los guiones (AAA-999[-99]); Circuito pasa de varchar(12) a varchar(14) en 5 tablas y se actualiza CK_Servicio_Circuito.", "sql/00_modelo/07", "José Hernández"],
    ["1.4", "20/10/2026", "4", "ctl.Reconciliacion y ctl.ReconciliacionPendiente (validación origen-destino); índice IX_ClienteOrigen_Sistema_Llave para las cargas de Books.", "sql/01_control/01; sql/00_modelo/08", "—"],
    ["1.5", "26/10/2026", "4", "core.ConsumoVoIP.RutaArchivo pasa de nvarchar(1000) a nvarchar(max) para conservar entero el HTML de las facturas VoIP; recarga de las 20.955 facturas recortadas.", "sql/00_modelo/09 y 10", "Josué Pérez"],
    ["2.0", "26/10/2026", "5", "Documento final: diccionario regenerado desde el código, descripción en todas las columnas, columna de disponibilidad para ask.data.cr / RAG, catálogo de reglas, historial de cambios y resolución de pendientes.", "herramientas/gen_diccionario.py", "Pendiente: José Hernández"],
    ["Matriz", "12/10/2026", "2", "16 filas de la matriz cambian respecto de la v1.0: Minutos_Customer y MinutosLegalId pasan de descartadas a mapeadas (D11); moneda de llamadas según D10; "
     "Telefon_a y VoIP_B2C a TieneVoIP; customer_name de conexiones a ContactoBooksId; duraciones de Navsat convertidas a segundos y minutos.", "etl/adn_etl/datos/matriz_correspondencia.csv", "José Hernández"],
]

PENDIENTES = [
    ["P-01", "Minutos_Customer y MinutosLegalId", "Se proponen como descartadas (tablas manuales de dic-2022). Confirmar que no se usan.", "José Hernández",
     "Decisión 1B (12/10): se guardan como histórico en core.MinutosContratadosHistorico (D11).", "Resuelto"],
    ["P-02", "Moneda de CallCost en llamadas", "Confirmar si CallCost está en USD (usa usd_exchange) o en colones, para llenar Moneda y CostoCRC.", "José Hernández",
     "Decisión 2C (12/10): destinos 506 en colones, el resto en dólares (D10).", "Resuelto"],
    ["P-03", "Tickets de Trengo sin ID_ZCRM", "Solo se vinculan si custom_data trae el ID del CRM; el teléfono no se usa (ID-05). Los demás quedan sin correspondencia.", "José Hernández",
     "Decisión 3 (12/10): histórico, sin excepción por teléfono (D12). Sprint 4: ninguno de los 22.560 tickets trae el ID del CRM; quedan sin vínculo.", "Resuelto"],
    ["P-04", "Conexiones/desconexiones", "Solo traen el nombre del cliente; por ID-05 quedan sin correspondencia salvo que exista otra llave.", "José Hernández",
     "customer_name resultó ser el ID del contacto de Books (5.350/5.350): se vinculan por esa llave (decisión 4, D12; corrección 03).", "Resuelto"],
    ["P-05", "Datos de tarjeta y credenciales en el CRM", "Accounts guarda Numero_Tarjeta y Accounts_Extended Passwordp/Userp. Se excluyen (IN-04); se recomienda revisarlos en el origen.", "José Hernández",
     "Decisión 5A (12/10): solo documentado; no se copian al modelo (IN-04). La revisión en el origen queda como recomendación del informe final.", "Documentado"],
    ["P-06", "Retención (D7)", "Confirmar requisitos legales y contables de retención de llamadas y facturación.", "José Hernández",
     "Decisión 6A (12/10): se mantiene D7 (24 meses en línea + histórico). El procedimiento de carga de 24 meses está en el manual de mantenimiento.", "Resuelto"],
    ["P-07", "Atributos extendidos de Zoho", "1.107 columnas quedan en JSON. Si alguna se usa en reportes, se promueve a columna propia en la versión 1.1.", "Responsables funcionales",
     "Decisión 7A (12/10): se dejan en JSON y se promueven según necesidad (procedimiento en el manual de mantenimiento). Para RAG quedan fuera hasta revisarlas.", "Resuelto"],
]


def vistas_rag(salida):
    """Vistas del esquema rag con solo las columnas Disponibles (lo que puede leer ask.data.cr)."""
    lineas = [
        "/* =====================================================================",
        "   Vistas para ask.data.cr y futuros sistemas RAG (Sprint 5) - GENERADO por herramientas/gen_diccionario.py",
        "   Aplica la decisión D6: cada vista expone solo las columnas clasificadas como 'Disponible' en el diccionario",
        "   (hoja Diccionario, columna 'ask.data.cr / RAG'). Las tablas 'No disponible' no tienen vista.",
        "   El rol rag_lector solo puede leer el esquema rag; no ve core, ref ni ctl.",
        "   Idempotente. SOLO sandbox: en producción requiere autorización de José Hernández.",
        "   ===================================================================== */",
        "SET QUOTED_IDENTIFIER ON; SET ANSI_NULLS ON;",
        "USE ADN_Unificado;", "GO",
        "IF SCHEMA_ID(N'rag') IS NULL EXEC(N'CREATE SCHEMA rag');", "GO",
    ]
    n = 0
    todas = [(t["schema"], t["name"], list(t["cols"]) + (list(LINAJE) if t["linaje"] else [])) for t in T]
    for schema, nombre, cols in todas:
        if rag.tabla(schema, nombre, [c["name"] for c in cols])[0] == "No disponible":
            continue
        visibles = [c["name"] for c in cols if rag.columna(schema, nombre, c["name"])[0] == rag.SI]
        if not visibles:
            continue
        n += 1
        lista = ",\n       ".join(f"[{c}]" for c in visibles)
        lineas += [f"CREATE OR ALTER VIEW rag.[{nombre}] AS", f"SELECT {lista}", f"FROM {schema}.[{nombre}];", "GO"]
    lineas += [
        "IF DATABASE_PRINCIPAL_ID(N'rag_lector') IS NULL CREATE ROLE rag_lector;", "GO",
        "GRANT SELECT ON SCHEMA::rag TO rag_lector;", "GO",
        "/* Verificación: vistas creadas y columnas expuestas */",
        "SELECT COUNT(DISTINCT v.object_id) AS Vistas, COUNT(*) AS Columnas",
        "FROM sys.views v JOIN sys.columns c ON c.object_id = v.object_id WHERE SCHEMA_NAME(v.schema_id) = N'rag';", "GO",
    ]
    os.makedirs(os.path.dirname(salida), exist_ok=True)
    with open(salida, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lineas) + "\n")
    return n


if __name__ == "__main__":
    salida = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, "docs", "Diccionario_Matriz_ADN_Unificado.xlsx")
    print(construir(salida), "→", salida)
    print("vistas rag:", vistas_rag(os.path.join(RAIZ, "sql", "30_rag", "vistas_rag.sql")))
